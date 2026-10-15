"""The month's pay as an Excel workbook: a summary per cutter and every record in detail.

Cells hold numbers and dates (not text), so the sheet can be sorted, filtered and summed.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CutRecord
from .reports import month_summary

NUMBER = '#,##0'
DATE = 'yyyy-mm-dd'
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1E3A5F')
TOTAL_FILL = PatternFill('solid', fgColor='EAEEF2')


def header_row(sheet, row, labels):
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def set_widths(sheet, widths):
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def summary_sheet(sheet, summary, T, month_label):
    sheet.title = T['xl_summary_sheet']
    sheet['A1'] = T['xl_title'].format(month=month_label)
    sheet['A1'].font = Font(bold=True, size=14, color='1E3A5F')
    sheet['A2'] = T['xl_period'].format(first=summary['first'].isoformat(), last=summary['last'].isoformat())

    labels = [T['cutter']]
    for cell in summary['style_cells']:
        labels += [T['xl_boxes'].format(style=cell['label']), T['xl_amount'].format(style=cell['label'])]
    labels += [T['xl_total_boxes'], T['xl_total_amount']]
    header_row(sheet, 4, labels)

    def write(row, name, cells, boxes, total):
        values = [name]
        for cell in cells:
            values += [cell['boxes'], int(cell['amount'])]
        values += [boxes, int(total)]
        for column, value in enumerate(values, start=1):
            target = sheet.cell(row=row, column=column, value=value)
            if column > 1:
                target.number_format = NUMBER

    row = 5
    for entry in summary['rows']:
        write(row, entry['cutter'].name, entry['style_cells'], entry['boxes'], entry['total'])
        row += 1
    write(row, T['xl_total_row'], summary['style_cells'], summary['boxes'], summary['total'])
    for column in range(1, len(labels) + 1):
        sheet.cell(row=row, column=column).font = Font(bold=True)
        sheet.cell(row=row, column=column).fill = TOTAL_FILL

    sheet.freeze_panes = 'B5'
    set_widths(sheet, [22] + [14] * (len(labels) - 1))


def detail_sheet(sheet, summary, T):
    labels = [T['date_col'], T['cutter'], T['style'], T['size'], T['boxes'], T['xl_rate'], T['xl_amount_col']]
    header_row(sheet, 1, labels)
    records = (
        CutRecord.objects.filter(work_date__range=(summary['first'], summary['last']))
        .select_related('cutter').order_by('work_date', 'cutter__name', 'batch_id', 'pk')
    )
    for row, record in enumerate(records, start=2):
        values = [record.work_date, record.cutter.name, record.get_style_display(), record.size,
                  record.boxes, int(record.unit_price), int(record.total)]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            if column == 1:
                cell.number_format = DATE
            elif column >= 5:
                cell.number_format = NUMBER
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(labels))}{max(sheet.max_row, 1)}'
    set_widths(sheet, [12, 20, 10, 10, 9, 12, 14])


def month_workbook_bytes(first_day, T):
    """The workbook for the month starting on `first_day`, as .xlsx bytes."""
    summary = month_summary(first_day)
    month_label = first_day.strftime('%Y-%m')
    workbook = Workbook()
    summary_sheet(workbook.active, summary, T, month_label)
    detail_sheet(workbook.create_sheet(T['xl_detail_sheet']), summary, T)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
