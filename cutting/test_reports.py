import uuid
from datetime import date
from io import BytesIO

from django.db.models import Sum
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape
from openpyxl import load_workbook

from . import reports, services
from .i18n import STRINGS
from .models import CutRecord, Cutter, PriceItem, PriceRate
from .tests import ManagerTestCase, make_user

EN, KO = STRINGS['en'], STRINGS['ko']
SEPT = date(2026, 9, 1)


def line(style, size):
    return PriceItem.objects.get(style=style, size=size)


class ReportData:
    """September 2026 for three cutters, with work just outside the month on both sides.

    통포 7 rises from ₩4,000 to ₩4,500 on the 16th, so Kim's month has two 통포 7 lines.
      Kim:  9/1 통포 7 × 10 (40,000) + 통할복 9 × 4 (10,000); 9/15 통포 7 × 5 (20,000); 9/30 통포 7 × 2 (9,000)
      Lee:  9/1 통포 12 × 3 (21,000)
      Park: 9/30 통할복 9 × 6 (15,000) — deactivated afterwards
      Outside: Kim 8/31 and 10/1, which September must not include.
    """

    def build(self):
        self.kim = Cutter.objects.create(name='Kim')
        self.lee = Cutter.objects.create(name='Lee')
        self.park = Cutter.objects.create(name='Park')
        self.tongpo7, self.tongpo12, self.halbok9 = line('tongpo', '7'), line('tongpo', '12'), line('tonghalbok', '9')
        PriceRate.objects.create(item=self.tongpo7, unit_price=4500, effective_from=date(2026, 9, 16))
        manager = getattr(self, 'manager', None)

        def save(day, *rows):
            services.save_batch(work_date=day, rows=list(rows), entered_by=manager, token=uuid.uuid4())

        save(date(2026, 9, 1), (self.kim, self.tongpo7, 10), (self.lee, self.tongpo12, 3))
        save(date(2026, 9, 1), (self.kim, self.halbok9, 4))  # a second batch the same day (style switch)
        save(date(2026, 9, 15), (self.kim, self.tongpo7, 5))
        save(date(2026, 9, 30), (self.kim, self.tongpo7, 2), (self.park, self.halbok9, 6))
        save(date(2026, 8, 31), (self.kim, self.tongpo7, 99))
        save(date(2026, 10, 1), (self.kim, self.tongpo7, 99))
        Cutter.objects.filter(pk=self.park.pk).update(is_active=False)
        self.park.refresh_from_db()


# ─── The arithmetic ───────────────────────────────────────────

class ReportMathTests(ReportData, TestCase):
    def setUp(self):
        self.build()

    def test_month_bounds(self):
        self.assertEqual(reports.month_bounds(date(2026, 9, 1)), (date(2026, 9, 1), date(2026, 9, 30)))
        self.assertEqual(reports.month_bounds(date(2026, 12, 1))[1], date(2026, 12, 31))
        self.assertEqual(reports.month_bounds(date(2028, 2, 1))[1], date(2028, 2, 29))  # leap year

    def test_a_month_is_provisional_until_its_last_day_has_passed(self):
        status = reports.month_status
        self.assertEqual(status(SEPT, today=date(2026, 9, 10)), {'provisional': True, 'through': date(2026, 9, 10)})
        self.assertEqual(status(SEPT, today=date(2026, 9, 30)), {'provisional': True, 'through': date(2026, 9, 30)})
        self.assertEqual(status(SEPT, today=date(2026, 10, 1)), {'provisional': False, 'through': date(2026, 9, 30)})
        self.assertEqual(reports.provisional_note(EN, SEPT, today=date(2026, 9, 10)),
                         EN['provisional_note'].format(date='2026-09-10'))
        self.assertEqual(reports.provisional_note(EN, SEPT, today=date(2026, 10, 1)), '')

    def test_month_summary_per_cutter_split_by_style(self):
        summary = reports.month_summary(SEPT)
        rows = {row['cutter'].name: row for row in summary['rows']}
        self.assertEqual(list(rows), ['Kim', 'Lee', 'Park'])  # inactive Park still gets paid
        kim = {cell['style']: (cell['boxes'], cell['amount']) for cell in rows['Kim']['style_cells']}
        self.assertEqual(kim, {'tongpo': (17, 69000), 'tonghalbok': (4, 10000)})
        self.assertEqual((rows['Kim']['boxes'], rows['Kim']['total']), (21, 79000))
        self.assertEqual((rows['Lee']['boxes'], rows['Lee']['total']), (3, 21000))
        self.assertEqual((rows['Park']['boxes'], rows['Park']['total']), (6, 15000))

    def test_month_totals_reconcile_with_the_records_and_exclude_other_months(self):
        summary = reports.month_summary(SEPT)
        recorded = CutRecord.objects.filter(work_date__range=(date(2026, 9, 1), date(2026, 9, 30)))
        self.assertEqual(summary['total'], recorded.aggregate(total=Sum('total'))['total'])
        self.assertEqual(summary['total'], 115000)
        self.assertEqual(summary['boxes'], 30)
        self.assertEqual({c['style']: c['amount'] for c in summary['style_cells']}, {'tongpo': 90000, 'tonghalbok': 25000})

    def test_a_mid_month_rate_change_shows_as_two_lines(self):
        month = reports.cutter_month(self.kim, SEPT)
        tongpo7 = [(l['unit_price'], l['boxes'], l['amount']) for l in month['lines'] if l['size'] == '7']
        self.assertEqual(tongpo7, [(4000, 15, 60000), (4500, 2, 9000)])
        groups = {g['style']: (g['boxes'], g['amount']) for g in month['style_groups']}
        self.assertEqual(groups, {'tongpo': (17, 69000), 'tonghalbok': (4, 10000)})
        self.assertEqual(month['total'], 79000)

    def test_cutter_month_breaks_down_by_day(self):
        days = reports.cutter_month(self.kim, SEPT)['days']
        self.assertEqual([d['date'] for d in days], [date(2026, 9, 1), date(2026, 9, 15), date(2026, 9, 30)])
        self.assertEqual([d['total'] for d in days], [50000, 20000, 9000])
        self.assertEqual([(l['style'], l['size']) for l in days[0]['lines']], [('tongpo', '7'), ('tonghalbok', '9')])

    def test_day_summary_adds_up_every_batch_per_cutter(self):
        day = {entry['cutter'].name: entry for entry in reports.day_summary(date(2026, 9, 1))}
        self.assertEqual(list(day), ['Kim', 'Lee'])
        self.assertEqual((day['Kim']['boxes'], day['Kim']['total']), (14, 50000))
        self.assertEqual(len(day['Kim']['lines']), 2)

    def test_cutters_paid_in_the_month(self):
        self.assertEqual([c.name for c in reports.cutters_paid_in(SEPT)], ['Kim', 'Lee', 'Park'])
        self.assertEqual(reports.cutters_paid_in(date(2026, 7, 1)), [])


# ─── The pages ────────────────────────────────────────────────

class ReportPageTests(ReportData, ManagerTestCase):
    def setUp(self):
        super().setUp()
        self.build()

    def month_url(self, name='cutting:month', *extra):
        return reverse(name, args=[SEPT, *extra])

    def test_month_page(self):
        response = self.client.get(self.month_url())
        self.assertContains(response, '₩115,000')
        self.assertContains(response, '₩79,000')
        self.assertContains(response, reverse('cutting:payslip', args=[SEPT, self.park.pk]))
        self.assertEqual(response.context['next_month'], date(2026, 10, 1))
        self.assertEqual(response.context['previous_month'], date(2026, 8, 1))

    def test_this_month_has_no_next_month(self):
        response = self.client.get(reverse('cutting:month_current'))
        self.assertEqual(response.context['month'], timezone.localdate().replace(day=1))
        self.assertIsNone(response.context['next_month'])

    def test_go_to_month_box(self):
        self.assertRedirects(self.client.get(reverse('cutting:month_current'), {'month': '2026-09'}), self.month_url())

    def test_an_impossible_month_is_a_404(self):
        self.assertEqual(self.client.get('/cutting/month/2026-13/').status_code, 404)

    def test_a_month_without_work(self):
        response = self.client.get(reverse('cutting:month', args=[date(2026, 7, 1)]))
        self.assertContains(response, EN['no_month_work'])
        self.assertNotContains(response, EN['export_excel'])

    def test_cutter_page(self):
        response = self.client.get(self.month_url('cutting:cutter_month', self.kim.pk))
        self.assertContains(response, escape(EN['cutter_month_title'].format(name='Kim', month='2026-09')))
        self.assertContains(response, reverse('cutting:day', args=[date(2026, 9, 15)]))
        self.assertEqual(response.context['data']['total'], 79000)

    def test_day_page_shows_each_cutters_work_and_pay(self):
        response = self.client.get(reverse('cutting:day', args=[date(2026, 9, 1)]))
        self.assertContains(response, EN['by_cutter'])
        self.assertContains(response, '₩50,000')  # Kim's two batches together
        self.assertContains(response, '10 × ₩4,000')

    def test_payslip(self):
        response = self.client.get(self.month_url('cutting:payslip', self.kim.pk))
        self.assertContains(response, 'Kim')
        self.assertContains(response, '₩4,500')  # the second 통포 7 line
        self.assertContains(response, EN['subtotal'].format(style='통포'))
        self.assertContains(response, '₩79,000')
        self.assertContains(response, EN['received_by'])

    def test_payslip_is_a_receipt_for_one_person(self):
        Cutter.objects.filter(pk=self.kim.pk).update(note='Late twice this month')
        response = self.client.get(self.month_url('cutting:payslip', self.kim.pk))
        self.assertNotContains(response, 'Late twice')  # the roster note is for managers, not the slip
        self.assertContains(response, '<title>Kim — Cutting pay slip — 2026-09</title>', html=False)
        self.assertContains(response, f"{escape(EN['daily_detail'])} · <span class=\"as-typed\">Kim</span> · 2026-09", html=False)
        # The day-by-day tally: 9/15, 5 boxes, ₩20,000.
        self.assertContains(response, '<span class="d">09-15</span><span class="b">5</span><span class="a">₩20,000</span>', html=False)
        all_slips = self.client.get(self.month_url('cutting:payslips'))
        self.assertEqual(all_slips.context['title'], f"{EN['payslip_title']} — 2026-09")

    def test_a_running_month_is_marked_provisional(self):
        today = timezone.localdate()
        note = escape(EN['provisional_note'].format(date=today.isoformat()))
        this_month = today.replace(day=1)
        for url in (reverse('cutting:month_current'),
                    reverse('cutting:cutter_month', args=[this_month, self.kim.pk]),
                    reverse('cutting:payslip', args=[this_month, self.kim.pk])):
            self.assertContains(self.client.get(url), note, msg_prefix=url)
        sheet = load_workbook(BytesIO(self.client.get(reverse('cutting:month_export', args=[this_month])).content)).active
        self.assertEqual(sheet['A3'].value, EN['provisional_note'].format(date=today.isoformat()))

    def test_a_finished_month_is_not_provisional(self):
        stem = EN['provisional_note'].split('—')[0]
        for url in (self.month_url(), self.month_url('cutting:cutter_month', self.kim.pk),
                    self.month_url('cutting:payslip', self.kim.pk)):
            self.assertNotContains(self.client.get(url), stem, msg_prefix=url)
        self.assertIsNone(load_workbook(BytesIO(self.client.get(self.month_url('cutting:month_export')).content)).active['A3'].value)

    def test_payslip_in_korean(self):
        self.use_korean()
        response = self.client.get(self.month_url('cutting:payslip', self.kim.pk))
        self.assertContains(response, 'lang="ko"')
        self.assertContains(response, KO['total_pay'])
        self.assertContains(response, KO['subtotal'].format(style='통할복'))

    def test_all_payslips_for_the_month(self):
        response = self.client.get(self.month_url('cutting:payslips'))
        self.assertEqual([slip['cutter'].name for slip in response.context['slips']], ['Kim', 'Lee', 'Park'])
        self.assertContains(response, 'class="sheet"', count=3)
        empty = self.client.get(reverse('cutting:payslips', args=[date(2026, 7, 1)]))
        self.assertContains(empty, EN['no_payslips'])

    def test_excel_export(self):
        response = self.client.get(self.month_url('cutting:month_export'))
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('fish-cutting-2026-09.xlsx', response['Content-Disposition'])

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, [EN['xl_summary_sheet'], EN['xl_detail_sheet']])
        summary = workbook[EN['xl_summary_sheet']]
        header = [cell.value for cell in summary[4]]
        self.assertEqual(header[0], EN['cutter'])
        self.assertEqual(header[-1], EN['xl_total_amount'])
        rows = [[cell.value for cell in row] for row in summary.iter_rows(min_row=5)]
        self.assertEqual(rows[0], ['Kim', 17, 69000, 4, 10000, 21, 79000])
        self.assertEqual(rows[-1], [EN['xl_total_row'], 20, 90000, 10, 25000, 30, 115000])

        detail = workbook[EN['xl_detail_sheet']]
        self.assertEqual(detail.max_row - 1, 6)  # six September records, nothing from Aug 31 or Oct 1
        first = [cell.value for cell in detail[2]]
        self.assertEqual(first[0].date(), date(2026, 9, 1))  # a real date, not text
        self.assertEqual(first[1:], ['Kim', '통포', '7', 10, 4000, 40000])

    def test_excel_keeps_typed_names_as_plain_text(self):
        formula = Cutter.objects.create(name='=1+1')
        control = Cutter.objects.create(name='Oh\x01Min')
        services.save_batch(work_date=date(2026, 9, 2), rows=[(formula, self.tongpo7, 1), (control, self.tongpo7, 1)],
                            entered_by=None, token=uuid.uuid4())
        response = self.client.get(self.month_url('cutting:month_export'))
        self.assertEqual(response.status_code, 200)  # openpyxl refuses control characters outright
        workbook = load_workbook(BytesIO(response.content))
        for sheet, column in ((workbook[EN['xl_summary_sheet']], 1), (workbook[EN['xl_detail_sheet']], 2)):
            names = {row[column - 1].value: row[column - 1] for row in sheet.iter_rows(min_row=2)}
            self.assertEqual(names['=1+1'].data_type, 's', sheet.title)  # text, never a live formula
            self.assertIn('OhMin', names, sheet.title)

    def test_excel_export_in_korean(self):
        self.use_korean()
        workbook = load_workbook(BytesIO(self.client.get(self.month_url('cutting:month_export')).content))
        self.assertEqual(workbook.sheetnames, [KO['xl_summary_sheet'], KO['xl_detail_sheet']])
        self.assertEqual(workbook[KO['xl_summary_sheet']]['A4'].value, KO['cutter'])

    def test_entry_and_report_pages_in_korean(self):
        self.use_korean()
        pages = {
            reverse('cutting:entry'): KO['entry_title'],
            reverse('cutting:day', args=[date(2026, 9, 1)]): KO['day_title'].format(date='2026-09-01'),
            self.month_url(): KO['month_title'].format(month='2026-09'),
            self.month_url('cutting:cutter_month', self.kim.pk): KO['cutter_month_title'].format(name='Kim', month='2026-09'),
        }
        for url, heading in pages.items():
            response = self.client.get(url)
            self.assertContains(response, 'lang="ko"', msg_prefix=url)
            self.assertContains(response, escape(heading), msg_prefix=url)

    def test_hub_links_to_monthly_pay(self):
        self.assertContains(self.client.get(reverse('cutting:hub')), reverse('cutting:month_current'))


class ReportAccessTests(ReportData, TestCase):
    def setUp(self):
        self.build()

    def urls(self):
        return [
            reverse('cutting:month_current'),
            reverse('cutting:month', args=[SEPT]),
            reverse('cutting:month_export', args=[SEPT]),
            reverse('cutting:payslips', args=[SEPT]),
            reverse('cutting:cutter_month', args=[SEPT, self.kim.pk]),
            reverse('cutting:payslip', args=[SEPT, self.kim.pk]),
        ]

    def test_only_managers_see_pay(self):
        for url in self.urls():
            self.assertTrue(self.client.get(url)['Location'].startswith('/login/'), url)
        self.client.force_login(make_user('worker', is_employee=True))
        for url in self.urls():
            self.assertEqual(self.client.get(url).status_code, 403, url)
        self.client.force_login(make_user('root', is_superuser=True, is_staff=True))
        for url in self.urls():
            self.assertEqual(self.client.get(url).status_code, 200, url)
